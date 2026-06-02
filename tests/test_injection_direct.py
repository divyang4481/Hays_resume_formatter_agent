"""
Direct injection test — bypasses API/LLM entirely.
Loads the template, injects hardcoded work_experience + education values,
saves the output, and verifies the content was replaced.
"""
import sys, json
from io import BytesIO
from pathlib import Path

sys.path.insert(0, '.')
from src.shared.repository import repo
from src.shared.storage import object_store
from src.worker.agents.resume_formatter.injector import inject_data_into_docx
from docx import Document
from docx.oxml.ns import qn

# ── 1. Load template ────────────────────────────────────────────────────────
TEMPLATE_ID = 'd1a6d0c3-1a73-4a50-a060-4b5e772a90bd'
template = repo.get_template(TEMPLATE_ID)
if not template:
    print("ERROR: Template not found"); sys.exit(1)

template_bytes = object_store.get_bytes(template['object_key'])
print(f"[OK] Template loaded: {template['object_key']} ({len(template_bytes)} bytes)")

# ── 2. Load manifest fields ─────────────────────────────────────────────────
manifest = repo.get_manifest(TEMPLATE_ID)
if not manifest:
    print("ERROR: Manifest not found"); sys.exit(1)

fields = manifest.get('fields', [])
print(f"[OK] Manifest loaded: {len(fields)} fields")

# ── 3. Hardcoded extracted values ───────────────────────────────────────────
extracted = {
    "work_experience": [
        {
            "job_description_date": "MAY 2020 — PRESENT",
            "organisation": "Cognizant, Bangalore",
            "bullet_point_responsibilities": [
                "Delivering code that meets security requirements",
                "Led solution architecture for retail loyalty program"
            ]
        },
        {
            "job_description_date": "JAN 2017 — APR 2020",
            "organisation": "Infosys, Pune",
            "bullet_point_responsibilities": [
                "Developed microservices architecture for banking",
                "Managed team of 8 engineers"
            ]
        },
        {
            "job_description_date": "JUN 2015 — DEC 2016",
            "organisation": "TCS, Mumbai",
            "bullet_point_responsibilities": [
                "Built REST APIs using Spring Boot",
            ]
        },
        {
            "job_description_date": "JAN 2010 — MAY 2015",
            "organisation": "Wipro, Hyderabad",
            "bullet_point_responsibilities": [
                "Cloud migration projects on AWS"
            ]
        },
    ],
    "education": [
        {
            "institution_date": "M.Sc., Gujarat University — 2002 to 2004",
            "bullet_point_grades": []
        },
        {
            "institution_date": "B.Sc., Gujarat University — 1999 to 2002",
            "bullet_point_grades": []
        },
    ],
    "candidate_name": "TEST CANDIDATE",
    "skills": ["Python", "Java", "AWS", "Docker"],
}

# ── 4. Run injection ─────────────────────────────────────────────────────────
print("\n[Running inject_data_into_docx...]")
output_bytes = inject_data_into_docx(template_bytes, extracted, fields)
print(f"[OK] Injection complete: {len(output_bytes)} bytes output")

# ── 5. Save output ───────────────────────────────────────────────────────────
out_path = Path('test_injection_direct.docx')
out_path.write_bytes(output_bytes)
print(f"[OK] Saved to: {out_path}")

# ── 6. Verify — inspect body paragraphs ─────────────────────────────────────
doc = Document(BytesIO(output_bytes))
NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
body_el = doc._element.find(f'.//{{{NS}}}body')

print("\n=== BODY PARAGRAPHS (direct children) ===")
para_idx = 0
found_we = []
found_ed = []
in_we = False
in_ed = False

for child in body_el:
    tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
    if tag != 'p':
        if tag == 'tbl':
            print(f"  [TBL]")
        continue

    t_text = ''.join(t.text or '' for t in child.findall(f'.//{{{NS}}}t')).strip()
    instr   = ''.join(t.text or '' for t in child.findall(f'.//{{{NS}}}instrText')).strip()
    has_fc  = bool(child.findall(f'.//{{{NS}}}fldChar'))
    runs    = child.findall(f'.//{{{NS}}}r')

    if 'WORK EXPERIENCE' in t_text:
        in_we, in_ed = True, False
        print(f"\n  p[{para_idx:03d}] >>> WORK EXPERIENCE <<<")
    elif 'EDUCATION' in t_text:
        in_we, in_ed = False, True
        print(f"\n  p[{para_idx:03d}] >>> EDUCATION <<<")
    elif 'INTERESTS' in t_text:
        in_we, in_ed = False, False
        print(f"\n  p[{para_idx:03d}] >>> INTERESTS <<<")
    elif in_we or in_ed:
        display = t_text[:80] if t_text else (f'[MACRO:{instr[:50]}]' if instr else '[EMPTY]')
        marker = '  **' if t_text else '  !!'
        print(f"{marker} p[{para_idx:03d}] runs={len(runs)} fldChar={has_fc} | {display}")
        if in_we and t_text:
            found_we.append(t_text[:60])
        if in_ed and t_text:
            found_ed.append(t_text[:60])

    para_idx += 1

print(f"\n=== RESULT ===")
print(f"work_experience text found: {len(found_we)} paragraphs")
for s in found_we:
    print(f"  - {s}")
print(f"education text found: {len(found_ed)} paragraphs")
for s in found_ed:
    print(f"  - {s}")

if found_we and found_ed:
    print("\n[PASS] Both work_experience and education were injected successfully!")
else:
    print("\n[FAIL] Injection failed — placeholders were not replaced.")
    # Show what the macrobutton instrText actually contains
    print("\n=== MACROBUTTON instrText samples ===")
    doc2 = Document(BytesIO(template_bytes))
    body2 = doc2._element.find(f'.//{{{NS}}}body')
    shown = 0
    for child in body2:
        if child.tag.split('}')[-1] != 'p':
            continue
        for instr_el in child.findall(f'.//{{{NS}}}instrText'):
            txt = (instr_el.text or '').strip()
            if txt and shown < 6:
                print(f"  instrText: {repr(txt)}")
                shown += 1
